from pathlib import Path
import sys

sys.path.insert(0, "/private/tmp/agi-proposal-pkgs")

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
LETTERHEAD = Path("/Users/mac/Downloads/AGI_Letterhead_Blank_Body.png")
OUTPUT = ROOT / "deliverables" / "AGI_Multi_Branch_ERP_Proposal.pdf"
WORD_OUTPUT = ROOT / "deliverables" / "AGI_Multi_Branch_ERP_Proposal.docx"

PURPLE = colors.HexColor("#32156f")
MID_PURPLE = colors.HexColor("#55308a")
INK = colors.HexColor("#202124")
MUTED = colors.HexColor("#5f6368")
PALE_PURPLE = colors.HexColor("#f1edf8")
RULE = colors.HexColor("#d8d0e7")


class LetterheadDoc(BaseDocTemplate):
    def __init__(self, filename):
        super().__init__(
            filename,
            pagesize=A4,
            leftMargin=22 * mm,
            rightMargin=22 * mm,
            topMargin=47 * mm,
            bottomMargin=22 * mm,
            title="Multi-Branch ERP Upgrade & Implementation Proposal",
            author="AGI Innovations Private Limited",
        )
        frame = Frame(
            self.leftMargin,
            self.bottomMargin,
            self.width,
            self.height,
            id="proposal",
        )
        self.addPageTemplates(PageTemplate(id="letterhead", frames=[frame], onPage=draw_letterhead))


def draw_letterhead(canvas, doc):
    canvas.saveState()
    canvas.drawImage(str(LETTERHEAD), 0, 0, width=A4[0], height=A4[1], mask="auto")
    canvas.setStrokeColor(RULE)
    canvas.setLineWidth(0.35)
    canvas.line(doc.leftMargin, 18 * mm, A4[0] - doc.rightMargin, 18 * mm)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 7.2)
    canvas.drawString(doc.leftMargin, 13.5 * mm, "Confidential | Multi-Branch ERP Upgrade & Implementation Proposal")
    canvas.drawRightString(A4[0] - doc.rightMargin, 13.5 * mm, f"Page {doc.page}")
    canvas.restoreState()


def paragraph(text, style):
    return Paragraph(text, style)


def bullets(items, style):
    return [paragraph(f"<bullet>&bull;</bullet>{item}", style) for item in items]


def section(title, body, styles):
    items = [paragraph(title, styles["h2"]), Spacer(1, 2.5 * mm)]
    if isinstance(body, str):
        items.append(paragraph(body, styles["body"]))
    else:
        for entry in body:
            items.append(entry if not isinstance(entry, str) else paragraph(entry, styles["body"]))
            items.append(Spacer(1, 1.6 * mm))
    items.append(Spacer(1, 3.2 * mm))
    return items


def labeled_table(rows, widths):
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PURPLE),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8.1),
                ("FONTSIZE", (0, 1), (-1, -1), 8),
                ("LEADING", (0, 1), (-1, -1), 10),
                ("TEXTCOLOR", (0, 1), (-1, -1), INK),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, RULE),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PALE_PURPLE]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def build():
    OUTPUT.parent.mkdir(exist_ok=True)
    doc = LetterheadDoc(str(OUTPUT))
    sample = getSampleStyleSheet()
    styles = {
        "kicker": ParagraphStyle("kicker", parent=sample["Normal"], fontName="Helvetica-Bold", fontSize=8.5, leading=11, textColor=MID_PURPLE, alignment=TA_CENTER, spaceAfter=3 * mm),
        "title": ParagraphStyle("title", parent=sample["Title"], fontName="Helvetica-Bold", fontSize=23, leading=28, textColor=PURPLE, alignment=TA_CENTER, spaceAfter=4 * mm),
        "subtitle": ParagraphStyle("subtitle", parent=sample["Normal"], fontName="Helvetica", fontSize=11, leading=15, textColor=INK, alignment=TA_CENTER, spaceAfter=7 * mm),
        "h2": ParagraphStyle("h2", parent=sample["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=PURPLE, spaceBefore=0, spaceAfter=0),
        "h3": ParagraphStyle("h3", parent=sample["Heading3"], fontName="Helvetica-Bold", fontSize=9.3, leading=12, textColor=PURPLE, spaceBefore=1.5 * mm, spaceAfter=1 * mm),
        "body": ParagraphStyle("body", parent=sample["BodyText"], fontName="Helvetica", fontSize=8.7, leading=12.1, textColor=INK, alignment=TA_LEFT, spaceAfter=0),
        "bullet": ParagraphStyle("bullet", parent=sample["BodyText"], fontName="Helvetica", fontSize=8.6, leading=11.7, leftIndent=4.5 * mm, firstLineIndent=-3.3 * mm, bulletIndent=0, textColor=INK, spaceAfter=1.2 * mm),
        "small": ParagraphStyle("small", parent=sample["BodyText"], fontName="Helvetica", fontSize=7.6, leading=9.8, textColor=MUTED),
        "quote": ParagraphStyle("quote", parent=sample["BodyText"], fontName="Helvetica-Bold", fontSize=11, leading=15, textColor=PURPLE, alignment=TA_CENTER),
    }

    story = []
    story += [
        Spacer(1, 14 * mm),
        paragraph("CONFIDENTIAL BUSINESS PROPOSAL", styles["kicker"]),
        paragraph("Multi-Branch ERP Upgrade &amp; Implementation", styles["title"]),
        paragraph("A secure, branch-aware operating platform for sales, service, inventory, procurement and finance", styles["subtitle"]),
    ]
    cover_rows = [
        ["Prepared for", "Client Management Team"],
        ["Prepared by", "AGI Innovations Private Limited"],
        ["Proposal date", "25 September 2026"],
        ["Implementation window", "8 weeks from project kickoff"],
        ["Commercial offer", "PKR 1,200,000 (Twelve Lakh) excluding applicable taxes"],
    ]
    cover = Table(cover_rows, colWidths=[42 * mm, 112 * mm], hAlign="CENTER")
    cover.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), PALE_PURPLE),
        ("TEXTCOLOR", (0, 0), (0, -1), PURPLE),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.2),
        ("LEADING", (0, 0), (-1, -1), 13),
        ("GRID", (0, 0), (-1, -1), 0.4, RULE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story += [cover, Spacer(1, 12 * mm), paragraph("A practical, phased path from the current ERP prototype to a production-ready platform that lets each branch operate independently while management retains the visibility it needs.", styles["quote"]), Spacer(1, 8 * mm), paragraph("This proposal is valid for 15 days from the proposal date.", styles["small"]), PageBreak()]

    story += section("1. Executive Summary", "AGI Innovations proposes to upgrade the existing ERP workspace into a secure multi-branch business system for the Client. The solution will support an initial rollout across Islamabad, Lahore and Karachi, with a controlled model that keeps each branch's operational data separated while allowing authorised head-office users to access consolidated reports and oversight dashboards.", styles)
    story += section("Business Outcomes", bullets([
        "A single cloud-based ERP platform, accessible securely from authorised branch locations.",
        "Branch-level separation for customers, quotations, orders, inventory, service activity and finance records.",
        "Role-based access so branch teams see only the data and actions relevant to their assigned branch.",
        "Head-office visibility of branch performance, stock, receivables and operational activity through consolidated reporting.",
        "A repeatable onboarding model for future branches without rebuilding the application.",
    ], styles["bullet"]), styles)
    story += section("Why This Approach", "The current system already demonstrates the desired business workflows and user experience. The proposed engagement focuses investment on the production capabilities required for real operations: secure authentication, central database persistence, enforced permissions, auditability, branch-aware rules, migration, user acceptance testing and go-live support.", styles)
    story += [PageBreak()]

    story += section("2. Proposed Solution & Scope", "The implementation will deliver one ERP instance with an initial configuration for the Islamabad, Lahore and Karachi branches. Each transaction will be assigned to a branch and protected by branch-aware access rules. Authorised corporate users can view consolidated information; branch users will work within their own permitted data boundary.", styles)
    scope_rows = [
        ["Workstream", "Included deliverables"],
        ["Platform foundation", "Production deployment; central database; secure sign-in; user roles; branch assignment; backup approach; audit-ready activity logging."],
        ["Branch operations", "Branch-specific customers, quotations, sales orders, deliveries, invoices, receipts, stock locations, service tickets, equipment and maintenance activity."],
        ["Commercial & inventory", "CRM; opportunity pipeline; quotations; order-to-delivery workflow; inventory, warehouses, transfers, batches/serials and stock visibility."],
        ["Procurement & service", "Suppliers; purchase requests/orders; goods receipts; equipment records; engineer workflow; tickets; contracts and maintenance scheduling."],
        ["Finance & oversight", "Invoices, receipts, expenses, receivables, management dashboards, branch and consolidated reports, approvals and exports."],
        ["Implementation services", "Discovery workshops; data templates; configuration; initial data migration assistance; testing; user training; go-live support and handover."],
    ]
    story += [labeled_table(scope_rows, [43 * mm, 111 * mm]), Spacer(1, 4 * mm)]
    story += section("Branch Data & Access Model", bullets([
        "Every user is assigned a role and one or more permitted branches.",
        "Operational users are restricted to their own branch data by default; cross-branch access requires explicit authorisation.",
        "Stock, documents and operational reports are branch tagged to prevent accidental mixing of branch records.",
        "Head-office management roles can receive consolidated dashboards and reports without giving all users unrestricted access.",
        "The architecture will allow additional branches to be configured using the same controls and deployment model.",
    ], styles["bullet"]), styles)
    story += [PageBreak()]

    story += section("3. Delivery Plan - 8 Weeks", "Delivery will be managed through weekly checkpoints with a nominated Client project owner. Timelines depend on timely decisions, master-data availability and user participation in testing/training.", styles)
    plan_rows = [
        ["Phase", "Timing", "Key outcomes"],
        ["1. Discovery & blueprint", "Week 1", "Confirm branch model, roles, processes, data sources, reports, priorities and acceptance criteria."],
        ["2. Production foundation", "Weeks 2-3", "Configure cloud environment, database, authentication, role/branch rules, base modules and data templates."],
        ["3. Workflow upgrades", "Weeks 3-5", "Build and validate branch-aware sales, inventory, procurement, service and finance workflows; reporting and approvals."],
        ["4. Migration & UAT", "Weeks 6-7", "Import agreed master/starting data, conduct user acceptance testing, fix agreed defects and finalise training materials."],
        ["5. Go-live & stabilisation", "Week 8", "Train users, launch initial branches, monitor adoption, resolve priority launch issues and hand over operating guidance."],
    ]
    story += [labeled_table(plan_rows, [35 * mm, 27 * mm, 92 * mm]), Spacer(1, 4 * mm)]
    story += section("Client Responsibilities", bullets([
        "Nominate one project owner and key users for each branch.",
        "Provide approved process decisions, user lists and clean source data within the agreed timetable.",
        "Provide timely feedback and sign-off during design, UAT and go-live stages.",
        "Arrange end-user attendance for training and maintain internal ownership of operating procedures after handover.",
    ], styles["bullet"]), styles)
    story += section("Acceptance", "Each milestone is accepted when the agreed features for that milestone are demonstrated and the Client confirms completion or provides a consolidated list of material deviations within five business days. Defects that materially prevent the agreed workflow will be corrected before go-live.", styles)
    story += [PageBreak()]

    story += section("4. Commercial Proposal", "AGI Innovations offers a fixed-price implementation for the defined Phase 1 scope and the initial three branches. This commercial structure gives the Client cost certainty while preserving clear controls over future enhancements and additional sites.", styles)
    commercial_rows = [
        ["Item", "Amount (PKR)"],
        ["Multi-Branch ERP Upgrade & Implementation - Phase 1", "1,200,000"],
        ["Initial branches included", "Islamabad, Lahore and Karachi"],
        ["Total one-time implementation fee", "PKR 1,200,000"],
        ["Taxes", "Applicable taxes, if any, are excluded and will be charged as required by law."],
    ]
    commercial = labeled_table(commercial_rows, [103 * mm, 51 * mm])
    commercial.setStyle(TableStyle([("FONTNAME", (0, 3), (-1, 3), "Helvetica-Bold"), ("BACKGROUND", (0, 3), (-1, 3), colors.HexColor("#e5dff2"))]))
    story += [commercial, Spacer(1, 5 * mm)]
    payment_rows = [
        ["Milestone", "Share", "Amount (PKR)"],
        ["Project kickoff and discovery approval", "40%", "480,000"],
        ["UAT-ready solution and migration completion", "40%", "480,000"],
        ["Go-live and handover", "20%", "240,000"],
    ]
    story += [paragraph("Payment Schedule", styles["h3"]), labeled_table(payment_rows, [85 * mm, 27 * mm, 42 * mm]), Spacer(1, 4 * mm)]
    story += section("Included Support", "Thirty calendar days of remote post-go-live warranty support is included for defects in the delivered scope. User questions and operational guidance will be handled during this period through the nominated Client project owner.", styles)
    story += [PageBreak()]

    story += section("5. Assumptions & Exclusions", "To protect both parties from scope ambiguity, the following are outside the fixed Phase 1 fee unless separately agreed in writing.", styles)
    story += section("Assumptions", bullets([
        "The quote covers three initial branches: Islamabad, Lahore and Karachi, with standard branch-aware workflows and permissions.",
        "The Client provides structured, reasonably clean master/starting data in the agreed templates; extensive legacy-data cleansing is not included.",
        "The Client makes key users available for workshops, testing and training throughout the eight-week plan.",
        "The implementation uses standard cloud services suitable for the agreed production environment.",
    ], styles["bullet"]), styles)
    story += section("Exclusions / Separately Priced Items", bullets([
        "Additional branches after the initial three; indicative configuration and onboarding fee: PKR 125,000 per branch, subject to scope confirmation.",
        "Third-party subscriptions, cloud hosting, domain, SMS/WhatsApp/email charges, payment gateway charges and hardware purchases.",
        "FBR/e-invoicing, banking, biometric, GPS, accounting-system or other third-party integrations.",
        "Custom mobile applications, offline field-service applications, advanced BI, OCR, bespoke report formats and major new modules.",
        "Historical transaction migration beyond the agreed opening data and approved migration templates.",
        "Ongoing maintenance, enhancements or support after the included 30-day warranty period; these may be covered by a separate support agreement.",
    ], styles["bullet"]), styles)
    story += section("Change Control", "Requests outside the agreed scope will be documented with their impact on timeline and cost. Work will begin only after written approval by both parties. This ensures the two-month target remains achievable.", styles)
    story += [PageBreak()]

    story += section("6. Next Steps", bullets([
        "Confirm acceptance of this proposal and the initial three-branch scope.",
        "Nominate the Client project owner and branch key users.",
        "Approve the project kickoff date and first milestone invoice.",
        "Conduct the Week 1 discovery workshop and issue the final solution blueprint.",
    ], styles["bullet"]), styles)
    story += [Spacer(1, 4 * mm), paragraph("We appreciate the opportunity to support the Client's growth with a practical, scalable ERP platform. AGI Innovations is committed to a disciplined rollout, transparent communication and measurable operational value.", styles["body"]), Spacer(1, 12 * mm)]
    sign_rows = [
        ["For AGI Innovations Private Limited", "For the Client"],
        ["\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________", "\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________"],
    ]
    signatures = Table(sign_rows, colWidths=[77 * mm, 77 * mm], hAlign="LEFT")
    signatures.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), PALE_PURPLE),
        ("TEXTCOLOR", (0, 0), (-1, 0), PURPLE),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.7),
        ("LEADING", (0, 0), (-1, -1), 12),
        ("GRID", (0, 0), (-1, -1), 0.4, RULE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story += [signatures]
    doc.build(story)
    print(OUTPUT)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def word_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, "32156F")
        set_cell_text(cell, header, bold=True, color=(255, 255, 255), size=8)
        if widths:
            cell.width = widths[index]
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if widths:
                cells[index].width = widths[index]
            shade(cells[index], "F1EDF8" if len(table.rows) % 2 else "FFFFFF")
            set_cell_text(cells[index], value, size=8)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build_word():
    assets = OUTPUT.parent / "proposal_assets"
    assets.mkdir(exist_ok=True)
    header_path = assets / "agi_letterhead_header.png"
    footer_path = assets / "agi_letterhead_footer.png"
    image = Image.open(LETTERHEAD)
    image.crop((0, 0, image.width, 238)).save(header_path)
    image.crop((0, 1424, image.width, image.height)).save(footer_path)

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(4.8)
    section.bottom_margin = Cm(2.1)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.5)
    normal.font.color.rgb = RGBColor(32, 33, 36)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, color in [("Title", 22, (50, 21, 111)), ("Heading 1", 13, (50, 21, 111)), ("Heading 2", 10, (50, 21, 111))]:
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_p.paragraph_format.space_after = Pt(0)
    header_p.add_run().add_picture(str(header_path), width=Cm(21))
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_after = Pt(0)
    footer_p.add_run().add_picture(str(footer_path), width=Cm(21))

    def p(text="", style=None, center=False, bold=False, color=None, size=None):
        para = document.add_paragraph(style=style)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        run = para.add_run(text)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)
        if size:
            run.font.size = Pt(size)
        return para

    def bullet(text):
        para = document.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(2)
        para.add_run(text)
        return para

    p("CONFIDENTIAL BUSINESS PROPOSAL", center=True, bold=True, color=(85, 48, 138), size=8.5)
    p("Multi-Branch ERP Upgrade & Implementation", style="Title", center=True)
    p("A secure, branch-aware operating platform for sales, service, inventory, procurement and finance", center=True, size=10.5)
    word_table(document, ["Prepared for", "Client Management Team"], [
        ["Prepared by", "AGI Innovations Private Limited"],
        ["Proposal date", "25 September 2026"],
        ["Implementation window", "8 weeks from project kickoff"],
        ["Commercial offer", "PKR 1,200,000 (Twelve Lakh) excluding applicable taxes"],
    ], [Cm(4.2), Cm(11.2)])
    p("A practical, phased path from the current ERP prototype to a production-ready platform that lets each branch operate independently while management retains the visibility it needs.", center=True, bold=True, color=(50, 21, 111), size=10)
    p("This proposal is valid for 15 days from the proposal date.", size=7.5, color=(95, 99, 104))
    document.add_page_break()

    p("1. Executive Summary", style="Heading 1")
    p("AGI Innovations proposes to upgrade the existing ERP workspace into a secure multi-branch business system for the Client. The solution will support an initial rollout across Islamabad, Lahore and Karachi, with a controlled model that keeps each branch's operational data separated while allowing authorised head-office users to access consolidated reports and oversight dashboards.")
    p("Business Outcomes", style="Heading 2")
    for text in [
        "A single cloud-based ERP platform, accessible securely from authorised branch locations.",
        "Branch-level separation for customers, quotations, orders, inventory, service activity and finance records.",
        "Role-based access so branch teams see only the data and actions relevant to their assigned branch.",
        "Head-office visibility through consolidated dashboards and reports.",
        "A repeatable onboarding model for future branches without rebuilding the application.",
    ]:
        bullet(text)
    p("Why This Approach", style="Heading 2")
    p("The current system already demonstrates the desired business workflows and user experience. The proposed engagement focuses investment on the production capabilities required for real operations: secure authentication, central database persistence, enforced permissions, auditability, branch-aware rules, migration, testing and go-live support.")
    document.add_page_break()

    p("2. Proposed Solution & Scope", style="Heading 1")
    p("The implementation will deliver one ERP instance with an initial configuration for the Islamabad, Lahore and Karachi branches. Each transaction will be assigned to a branch and protected by branch-aware access rules. Authorised corporate users can view consolidated information; branch users will work within their own permitted data boundary.")
    word_table(document, ["Workstream", "Included deliverables"], [
        ["Platform foundation", "Production deployment; central database; secure sign-in; user roles; branch assignment; backup approach; audit-ready activity logging."],
        ["Branch operations", "Branch-specific customers, quotations, sales orders, deliveries, invoices, receipts, stock locations, service tickets, equipment and maintenance activity."],
        ["Commercial & inventory", "CRM; opportunities; quotations; order-to-delivery workflow; inventory, warehouses, transfers, batches/serials and stock visibility."],
        ["Procurement & service", "Suppliers; purchase requests/orders; goods receipts; equipment records; engineers; tickets; contracts and maintenance scheduling."],
        ["Finance & oversight", "Invoices, receipts, expenses, receivables, management dashboards, branch and consolidated reports, approvals and exports."],
        ["Implementation services", "Discovery workshops; data templates; configuration; initial data migration assistance; testing; user training; go-live support and handover."],
    ], [Cm(4.2), Cm(11.2)])
    p("Branch Data & Access Model", style="Heading 2")
    for text in [
        "Every user is assigned a role and one or more permitted branches.",
        "Operational users are restricted to their own branch data by default; cross-branch access requires explicit authorisation.",
        "Stock, documents and operational reports are branch tagged to prevent accidental mixing of branch records.",
        "Head-office management can receive consolidated dashboards and reports without giving all users unrestricted access.",
    ]:
        bullet(text)
    document.add_page_break()

    p("3. Delivery Plan - 8 Weeks", style="Heading 1")
    p("Delivery will be managed through weekly checkpoints with a nominated Client project owner. Timelines depend on timely decisions, master-data availability and user participation in testing/training.")
    word_table(document, ["Phase", "Timing", "Key outcomes"], [
        ["Discovery & blueprint", "Week 1", "Confirm branch model, roles, processes, data sources, reports, priorities and acceptance criteria."],
        ["Production foundation", "Weeks 2-3", "Configure cloud environment, database, authentication, role/branch rules, base modules and data templates."],
        ["Workflow upgrades", "Weeks 3-5", "Build and validate branch-aware sales, inventory, procurement, service and finance workflows; reporting and approvals."],
        ["Migration & UAT", "Weeks 6-7", "Import agreed master/starting data, conduct UAT, fix agreed defects and finalise training materials."],
        ["Go-live & stabilisation", "Week 8", "Train users, launch initial branches, monitor adoption, resolve priority launch issues and hand over guidance."],
    ], [Cm(3.4), Cm(2.6), Cm(9.4)])
    p("Client Responsibilities", style="Heading 2")
    for text in [
        "Nominate one project owner and key users for each branch.",
        "Provide approved process decisions, user lists and clean source data within the agreed timetable.",
        "Provide timely feedback and sign-off during design, UAT and go-live stages.",
        "Arrange end-user attendance for training and maintain internal ownership of procedures after handover.",
    ]:
        bullet(text)
    p("Acceptance", style="Heading 2")
    p("Each milestone is accepted when the agreed features are demonstrated and the Client confirms completion or provides a consolidated list of material deviations within five business days.")
    document.add_page_break()

    p("4. Commercial Proposal", style="Heading 1")
    p("AGI Innovations offers a fixed-price implementation for the defined Phase 1 scope and the initial three branches. This commercial structure gives the Client cost certainty while preserving clear controls over future enhancements and additional sites.")
    word_table(document, ["Item", "Amount (PKR)"], [
        ["Multi-Branch ERP Upgrade & Implementation - Phase 1", "1,200,000"],
        ["Initial branches included", "Islamabad, Lahore and Karachi"],
        ["Total one-time implementation fee", "PKR 1,200,000"],
        ["Taxes", "Applicable taxes, if any, are excluded and will be charged as required by law."],
    ], [Cm(10.5), Cm(5.0)])
    p("Payment Schedule", style="Heading 2")
    word_table(document, ["Milestone", "Share", "Amount (PKR)"], [
        ["Project kickoff and discovery approval", "40%", "480,000"],
        ["UAT-ready solution and migration completion", "40%", "480,000"],
        ["Go-live and handover", "20%", "240,000"],
    ], [Cm(8.5), Cm(2.7), Cm(4.3)])
    p("Included Support", style="Heading 2")
    p("Thirty calendar days of remote post-go-live warranty support is included for defects in the delivered scope. User questions and operational guidance will be handled through the nominated Client project owner.")
    document.add_page_break()

    p("5. Assumptions & Exclusions", style="Heading 1")
    p("To protect both parties from scope ambiguity, the following are outside the fixed Phase 1 fee unless separately agreed in writing.")
    p("Assumptions", style="Heading 2")
    for text in [
        "The quote covers three initial branches: Islamabad, Lahore and Karachi, with standard branch-aware workflows and permissions.",
        "The Client provides structured, reasonably clean master/starting data in the agreed templates; extensive legacy-data cleansing is not included.",
        "The Client makes key users available for workshops, testing and training throughout the eight-week plan.",
        "The implementation uses standard cloud services suitable for the agreed production environment.",
    ]:
        bullet(text)
    p("Exclusions / Separately Priced Items", style="Heading 2")
    for text in [
        "Additional branches after the initial three; indicative configuration and onboarding fee: PKR 125,000 per branch, subject to scope confirmation.",
        "Third-party subscriptions, cloud hosting, domain, SMS/WhatsApp/email charges, payment gateway charges and hardware purchases.",
        "FBR/e-invoicing, banking, biometric, GPS, accounting-system or other third-party integrations.",
        "Custom mobile applications, offline field-service applications, advanced BI, OCR, bespoke report formats and major new modules.",
        "Historical transaction migration beyond the agreed opening data and approved migration templates.",
        "Ongoing maintenance, enhancements or support after the included 30-day warranty period.",
    ]:
        bullet(text)
    p("Change Control", style="Heading 2")
    p("Requests outside the agreed scope will be documented with their impact on timeline and cost. Work will begin only after written approval by both parties. This ensures the two-month target remains achievable.")
    document.add_page_break()

    p("6. Next Steps", style="Heading 1")
    for text in [
        "Confirm acceptance of this proposal and the initial three-branch scope.",
        "Nominate the Client project owner and branch key users.",
        "Approve the project kickoff date and first milestone invoice.",
        "Conduct the Week 1 discovery workshop and issue the final solution blueprint.",
    ]:
        bullet(text)
    p("We appreciate the opportunity to support the Client's growth with a practical, scalable ERP platform. AGI Innovations is committed to a disciplined rollout, transparent communication and measurable operational value.")
    document.add_paragraph("\n\n")
    word_table(document, ["For AGI Innovations Private Limited", "For the Client"], [["\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________", "\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________"]], [Cm(7.75), Cm(7.75)])
    document.save(WORD_OUTPUT)
    print(WORD_OUTPUT)


if __name__ == "__main__":
    build()
    build_word()
