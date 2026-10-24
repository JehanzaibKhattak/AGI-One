from pathlib import Path
import sys

sys.path.insert(0, "/private/tmp/agi-proposal-pkgs")
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import BaseDocTemplate, Frame, PageBreak, PageTemplate, Spacer, Table, TableStyle

from generate_integrated_proposal import (
    LETTERHEAD, OUT_DIR, PDF_OUT, ProposalDoc, PURPLE, LILAC, RULE,
    bullets, p, pdf_styles, table,
)


PDF_OUTPUT = OUT_DIR / "AGI_AI_Intelligence_Module_Proposal.pdf"
DOCX_OUTPUT = OUT_DIR / "AGI_AI_Intelligence_Module_Proposal.docx"


def ai_letterhead(canvas, doc):
    canvas.saveState()
    canvas.drawImage(str(LETTERHEAD), 0, 0, A4[0], A4[1], mask="auto")
    canvas.setStrokeColor(RULE)
    canvas.setLineWidth(.35)
    canvas.line(doc.leftMargin, 18 * mm, A4[0] - doc.rightMargin, 18 * mm)
    canvas.setFillColor(colors.HexColor("#5F6368"))
    canvas.setFont("Helvetica", 7.2)
    canvas.drawString(doc.leftMargin, 13.5 * mm, "Confidential | AI Intelligence Module Proposal")
    canvas.drawRightString(A4[0] - doc.rightMargin, 13.5 * mm, f"Page {doc.page}")
    canvas.restoreState()


class AIProposalDoc(BaseDocTemplate):
    def __init__(self, filename):
        super().__init__(filename, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=47 * mm, bottomMargin=22 * mm)
        self.addPageTemplates(PageTemplate(id="ai", frames=[Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="f")], onPage=ai_letterhead))


def build_pdf():
    OUT_DIR.mkdir(exist_ok=True)
    s = pdf_styles()
    story = [Spacer(1, 13 * mm), p("CONFIDENTIAL BUSINESS PROPOSAL", s["kicker"]), p("AI Intelligence Module", s["title"]), p("A secure, source-grounded AI assistant for the Client's business management system", s["sub"])]
    story += [table(["Proposal detail", "Value"], [
        ["Prepared for", "Client Management Team"],
        ["Prepared by", "AGI Innovations Private Limited"],
        ["Proposal date", "25 September 2026"],
        ["Estimated delivery", "3-5 weeks from approved data readiness"],
        ["Implementation fee", "USD 2,500 / PKR equivalent at invoice date"],
    ], [45 * mm, 109 * mm]), Spacer(1, 9 * mm), p("Give authorised users faster access to approved business knowledge and permitted operational information - without compromising branch, role or data boundaries.", s["quote"]), Spacer(1, 5 * mm), p("This proposal is valid for 15 days from the proposal date.", s["note"]), PageBreak()]

    story += [p("1. Executive Summary", s["h1"]), p("AGI Innovations proposes a controlled AI Intelligence Module integrated into the Client's existing business management system. The assistant will retrieve answers from approved documents and permitted business records, allowing staff to search, understand and summarise information in natural language while preserving user and branch permissions.", s["body"]), p("What this solves", s["h2"])]
    story += bullets(["Staff spend less time searching policies, product information, procedures and operational records.", "Approved users receive concise, source-referenced answers in the context of the systems they already use.", "Branch and role controls remain in force: users can retrieve only information they are permitted to view.", "Management can improve knowledge consistency while retaining human review for business decisions."], s)
    story += [p("Important clarification", s["h2"]), p("This is not a new large language model trained from scratch. It is a retrieval-augmented generation (RAG) solution: the AI retrieves relevant approved sources at answer time and uses them to produce a grounded response with references.", s["body"]), p("Value posture", s["h2"]), p("The expected productivity and knowledge-access benefits are directional objectives. Actual time savings, adoption and business impact will be measured after launch using agreed user questions and usage feedback.", s["note"]), PageBreak()]

    story += [p("2. Included Scope", s["h1"]), table(["Workstream", "Included deliverables"], [
        ["AI assistant experience", "A natural-language chat experience integrated into the Client's business management system."],
        ["Knowledge ingestion", "Ingestion and indexing of approved documents, SOPs, policies, product information and agreed business records."],
        ["Permission-aware retrieval", "Role and branch filters applied before retrieval so users access only permitted content."],
        ["Grounded answers", "Source references with responses and a safe fallback when approved sources do not support an answer."],
        ["Operational insights", "Read-only search, summaries and permitted operational questions for the initial release."],
        ["Quality & handover", "Agreed test-question set, evaluation, administrator guidance, user orientation and controlled launch."],
    ], [48 * mm, 106 * mm]), Spacer(1, 3 * mm), p("Priority initial use cases", s["h2"])]
    story += bullets(["Ask questions across approved SOPs, policies, product information and internal knowledge.", "Summarise relevant permitted records, notes and operational information.", "Guide users toward the relevant process, record or source document.", "Provide management-approved knowledge answers consistently across authorised teams."], s)
    story += [PageBreak()]

    story += [p("3. Data Security & AI Guardrails", s["h1"]), p("The AI module is designed as an assistive capability, not an autonomous decision-maker. It will operate within agreed data, access and review controls.", s["body"]), table(["Control", "How it protects the Client"], [
        ["Approved source boundary", "Only sources approved by the Client's nominated data owner are indexed for AI use."],
        ["Role and branch controls", "Retrieval respects application permissions, including branch-level data separation where configured."],
        ["Source references", "Users can see the source material supporting an answer and verify it before acting."],
        ["Safe fallback", "The assistant can state that it cannot find support in approved material rather than inventing an answer."],
        ["Read-only initial release", "The assistant will not create, change, approve or post business transactions in the initial scope."],
        ["Human accountability", "Users and designated approvers remain responsible for commercial, financial and operational decisions."],
    ], [48 * mm, 106 * mm]), Spacer(1, 3 * mm), p("Data readiness", s["h2"]), p("The Client will nominate a data owner, provide approved and reasonably organised source material, and confirm the user/branch permission model. Extensive data cleansing, broad historical migration, complex multilingual corpus preparation and private/on-premise model infrastructure are outside this fixed scope unless separately agreed.", s["body"]), PageBreak()]

    story += [p("4. Delivery Plan", s["h1"]), p("The module can be delivered in approximately three to five weeks after the Client confirms data readiness and assigns the required owners.", s["body"]), table(["Stage", "Timing", "Outcome"], [
        ["Discovery & design", "Week 1", "Confirm sources, user groups, data boundaries, test questions, guardrails and acceptance criteria."],
        ["Build & integration", "Weeks 2-3", "Set up knowledge ingestion, retrieval, application integration, permission filtering and answer citations."],
        ["Evaluation & refinement", "Week 4", "Test against agreed real business questions, resolve material issues and finalise administrator guidance."],
        ["Controlled launch", "Week 5 if required", "Train users, release to approved groups, monitor early use and hand over operating guidance."],
    ], [45 * mm, 31 * mm, 78 * mm]), Spacer(1, 3 * mm), p("Client responsibilities", s["h2"])]
    story += bullets(["Nominate one project owner and one AI data owner.", "Provide approved source documents, user groups and access decisions on the agreed timetable.", "Supply and validate a practical set of real business questions for testing.", "Make authorised users available for acceptance testing and launch orientation."], s)
    story += [PageBreak()]

    story += [p("5. Commercial Proposal", s["h1"]), table(["Commercial item", "Value"], [
        ["AI Intelligence Module - one-time implementation", "USD 2,500 / PKR equivalent at invoice date"],
        ["Implementation period", "3-5 weeks from approved data readiness"],
        ["Included post-launch support", "30 calendar days of remote warranty support for delivered-scope defects"],
        ["Taxes", "Applicable taxes, if any, are excluded and charged as required by law"],
    ], [72 * mm, 82 * mm]), Spacer(1, 4 * mm), p("Payment schedule", s["h2"]), table(["Milestone", "Share", "Amount"], [
        ["Kickoff and AI discovery approval", "50%", "USD 1,250"],
        ["UAT-ready solution", "30%", "USD 750"],
        ["Controlled launch and handover", "20%", "USD 500"],
    ], [78 * mm, 29 * mm, 47 * mm]), Spacer(1, 3 * mm), p("Recurring services", s["h2"]), p("Model/API usage, cloud hosting, vector database, domain and third-party service charges are excluded from the one-time implementation fee and will be charged separately based on the selected service providers and actual usage.", s["body"]), p("Items requiring separate scope", s["h2"])]
    story += bullets(["Autonomous write-actions, approvals or transaction creation.", "Private/on-premise model hosting, advanced integrations and custom mobile AI applications.", "Extensive source-document cleansing, custom analytics, OCR and multilingual model training.", "Ongoing managed support, monitoring or enhancements after the included warranty period."], s)
    story += [PageBreak()]

    story += [p("6. Acceptance & Next Decision", s["h1"]), p("The implementation is accepted when the AI assistant demonstrates the agreed core use cases, observes the agreed permission controls, returns source-referenced responses where relevant, and completes the jointly approved evaluation questions. Material defects that prevent the agreed workflows will be corrected before controlled launch.", s["body"]), p("Recommended next decision", s["h2"]), p("Approve the AI Intelligence Module, nominate the project and data owners, and schedule the Week 1 discovery session to confirm sources, access controls and the evaluation question set.", s["body"]), Spacer(1, 13 * mm)]
    signatures = Table([["For AGI Innovations Private Limited", "For the Client"], ["\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________", "\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________"]], colWidths=[77 * mm, 77 * mm])
    signatures.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), LILAC), ("TEXTCOLOR", (0, 0), (-1, 0), PURPLE), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, 1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8.7), ("LEADING", (0, 0), (-1, -1), 12), ("GRID", (0, 0), (-1, -1), .4, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    story += [signatures]
    AIProposalDoc(str(PDF_OUTPUT)).build(story)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr(); element = OxmlElement("w:shd"); element.set(qn("w:fill"), fill); props.append(element)


def word_cell(cell, text, bold=False, white=False):
    cell.text = ""; para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text); run.font.name = "Arial"; run.font.size = Pt(8); run.bold = bold
    run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(32, 33, 36); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def word_table(doc, headers, rows):
    result = doc.add_table(rows=1, cols=len(headers)); result.alignment = WD_TABLE_ALIGNMENT.CENTER; result.style = "Table Grid"
    for i, header in enumerate(headers): shade(result.rows[0].cells[i], "32156F"); word_cell(result.rows[0].cells[i], header, True, True)
    for index, row in enumerate(rows):
        cells = result.add_row().cells
        for i, value in enumerate(row): shade(cells[i], "F1EDF8" if index % 2 else "FFFFFF"); word_cell(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_docx():
    assets = OUT_DIR / "proposal_assets"; assets.mkdir(exist_ok=True)
    head, foot = assets / "agi_ai_header.png", assets / "agi_ai_footer.png"
    image = Image.open(LETTERHEAD); image.crop((0, 0, image.width, 238)).save(head); image.crop((0, 1424, image.width, image.height)).save(foot)
    doc = Document(); section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7); section.left_margin, section.right_margin, section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.2), Cm(4.8), Cm(2.1); section.header_distance, section.footer_distance = Cm(0), Cm(0)
    for name, size, color in [("Normal", 8.5, (32, 33, 36)), ("Title", 21, (50, 21, 111)), ("Heading 1", 13, (50, 21, 111)), ("Heading 2", 10, (50, 21, 111))]:
        style = doc.styles[name]; style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor(*color)
        if name != "Normal": style.font.bold = True
    h = section.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.add_run().add_picture(str(head), width=Cm(21))
    f = section.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER; f.add_run().add_picture(str(foot), width=Cm(21))
    def para(text="", style=None, center=False, bold=False, color=None, size=None):
        q = doc.add_paragraph(style=style); q.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT; r = q.add_run(text); r.bold = bold
        if color: r.font.color.rgb = RGBColor(*color)
        if size: r.font.size = Pt(size)
        return q
    def item(text):
        q = doc.add_paragraph(style="List Bullet"); q.paragraph_format.space_after = Pt(2); q.add_run(text)
    para("CONFIDENTIAL BUSINESS PROPOSAL", center=True, bold=True, color=(85, 48, 138), size=8.5); para("AI Intelligence Module", "Title", True); para("A secure, source-grounded AI assistant for the Client's business management system", center=True, size=10.5)
    word_table(doc, ["Proposal detail", "Value"], [["Prepared for", "Client Management Team"], ["Prepared by", "AGI Innovations Private Limited"], ["Proposal date", "25 September 2026"], ["Estimated delivery", "3-5 weeks from approved data readiness"], ["Implementation fee", "USD 2,500 / PKR equivalent at invoice date"]])
    para("Give authorised users faster access to approved business knowledge and permitted operational information - without compromising branch, role or data boundaries.", center=True, bold=True, color=(50, 21, 111), size=10); doc.add_page_break()
    sections = [
        ("1. Executive Summary", "AGI Innovations proposes a controlled AI Intelligence Module integrated into the Client's existing business management system. The assistant retrieves answers from approved documents and permitted business records, allowing natural-language search and summaries while preserving user and branch permissions.", ["This is a RAG solution, not a new model trained from scratch.", "Users receive source-referenced answers from approved material.", "Branch and role controls remain in force.", "The module assists users; it does not replace human decision-making."]),
        ("2. Included Scope", "The v1 module includes an integrated chat experience, ingestion/indexing of approved documents and agreed records, role/branch-aware retrieval, source citations, safe fallback, read-only insights, testing and administrator handover.", ["Ask questions across approved SOPs, policies, product information and internal knowledge.", "Summarise relevant permitted records and operational information.", "Guide users to the relevant process, record or source document.", "Test against real Client questions before controlled launch."]),
        ("3. Data Security & Guardrails", "Only sources approved by the Client's nominated data owner will be indexed. Retrieval follows existing permissions. The initial release will not create, change, approve or post business transactions.", ["Source references help users verify responses.", "The assistant can decline when approved sources do not support an answer.", "Human users and designated approvers remain responsible for business decisions.", "Extensive data cleansing and private/on-premise model infrastructure require separate scope."]),
        ("4. Delivery Plan", "Estimated delivery is three to five weeks: Week 1 discovery and source/permission design; Weeks 2-3 build and integration; Week 4 evaluation and refinement; Week 5 controlled launch if required.", ["The Client nominates a project owner and AI data owner.", "The Client provides approved source documents, user groups and access decisions.", "The Client validates a practical set of real business questions for UAT."]),
    ]
    for title, text, points in sections:
        para(title, "Heading 1"); para(text)
        for point in points: item(point)
        doc.add_page_break()
    para("5. Commercial Proposal", "Heading 1")
    word_table(doc, ["Commercial item", "Value"], [["AI Intelligence Module - one-time implementation", "USD 2,500 / PKR equivalent at invoice date"], ["Implementation period", "3-5 weeks from approved data readiness"], ["Included post-launch support", "30 calendar days of remote warranty support for delivered-scope defects"], ["Taxes", "Excluded; charged as required by law"]])
    para("Payment Schedule", "Heading 2"); word_table(doc, ["Milestone", "Share", "Amount"], [["Kickoff and AI discovery", "50%", "USD 1,250"], ["UAT-ready solution", "30%", "USD 750"], ["Controlled launch and handover", "20%", "USD 500"]])
    para("Recurring services", "Heading 2"); para("Model/API usage, cloud hosting, vector database, domain and third-party service charges are excluded and charged separately based on selected providers and actual usage.")
    para("6. Acceptance & Next Decision", "Heading 1"); para("The solution is accepted when it demonstrates the agreed core use cases, observes permission controls, returns source-referenced responses where relevant, and completes the jointly approved evaluation questions.")
    para("Recommended Decision", "Heading 2"); para("Approve the AI Intelligence Module, nominate the project and data owners, and schedule the discovery session to confirm sources, access controls and the evaluation question set.")
    word_table(doc, ["For AGI Innovations Private Limited", "For the Client"], [["\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________", "\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________"]])
    doc.save(DOCX_OUTPUT)


if __name__ == "__main__":
    build_pdf(); build_docx(); print(PDF_OUTPUT); print(DOCX_OUTPUT)
