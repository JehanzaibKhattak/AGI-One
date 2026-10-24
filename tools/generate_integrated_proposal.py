from pathlib import Path
import sys

sys.path.insert(0, "/private/tmp/agi-proposal-pkgs")

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
LETTERHEAD = Path("/Users/mac/Downloads/AGI_Letterhead_Blank_Body.png")
OUT_DIR = ROOT / "deliverables"
PDF_OUT = OUT_DIR / "AGI_Integrated_ERP_and_AI_Proposal.pdf"
DOCX_OUT = OUT_DIR / "AGI_Integrated_ERP_and_AI_Proposal.docx"

PURPLE = colors.HexColor("#32156F")
MID = colors.HexColor("#55308A")
INK = colors.HexColor("#202124")
MUTED = colors.HexColor("#5F6368")
LILAC = colors.HexColor("#F1EDF8")
RULE = colors.HexColor("#D8D0E7")


class ProposalDoc(BaseDocTemplate):
    def __init__(self, filename):
        super().__init__(filename, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=47 * mm, bottomMargin=22 * mm)
        self.addPageTemplates(PageTemplate(id="agi", frames=[Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="f")], onPage=letterhead))


def letterhead(canvas, doc):
    canvas.saveState()
    canvas.drawImage(str(LETTERHEAD), 0, 0, A4[0], A4[1], mask="auto")
    canvas.setStrokeColor(RULE)
    canvas.setLineWidth(.35)
    canvas.line(doc.leftMargin, 18 * mm, A4[0] - doc.rightMargin, 18 * mm)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 7.2)
    canvas.drawString(doc.leftMargin, 13.5 * mm, "Confidential | Integrated ERP & AI Business Management Proposal")
    canvas.drawRightString(A4[0] - doc.rightMargin, 13.5 * mm, f"Page {doc.page}")
    canvas.restoreState()


def pdf_styles():
    base = getSampleStyleSheet()
    return {
        "kicker": ParagraphStyle("k", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=8.5, leading=11, textColor=MID, alignment=TA_CENTER, spaceAfter=3 * mm),
        "title": ParagraphStyle("t", parent=base["Title"], fontName="Helvetica-Bold", fontSize=22, leading=27, textColor=PURPLE, alignment=TA_CENTER, spaceAfter=4 * mm),
        "sub": ParagraphStyle("s", parent=base["Normal"], fontName="Helvetica", fontSize=10.5, leading=14, textColor=INK, alignment=TA_CENTER, spaceAfter=6 * mm),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=PURPLE, spaceAfter=2.5 * mm),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=9.5, leading=12, textColor=PURPLE, spaceBefore=1.5 * mm, spaceAfter=1.2 * mm),
        "body": ParagraphStyle("b", parent=base["BodyText"], fontName="Helvetica", fontSize=8.7, leading=12, textColor=INK, spaceAfter=1.8 * mm),
        "bullet": ParagraphStyle("bl", parent=base["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11.5, textColor=INK, leftIndent=4.6 * mm, firstLineIndent=-3.3 * mm, spaceAfter=1.2 * mm),
        "note": ParagraphStyle("n", parent=base["BodyText"], fontName="Helvetica", fontSize=7.6, leading=9.5, textColor=MUTED),
        "quote": ParagraphStyle("q", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=11, leading=15, textColor=PURPLE, alignment=TA_CENTER),
    }


def p(text, style):
    return Paragraph(text, style)


def bullets(items, styles):
    return [p(f"<bullet>&bull;</bullet>{item}", styles["bullet"]) for item in items]


def table(headers, rows, widths):
    data = [[p(h, ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white, leading=10)) for h in headers]]
    data += [[p(str(cell), ParagraphStyle("td", fontName="Helvetica", fontSize=8, textColor=INK, leading=10.2)) for cell in row] for row in rows]
    result = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    result.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), PURPLE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), .35, RULE), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LILAC]),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return result


def build_pdf():
    OUT_DIR.mkdir(exist_ok=True)
    styles = pdf_styles()
    story = [Spacer(1, 12 * mm), p("CONFIDENTIAL BUSINESS PROPOSAL", styles["kicker"]), p("Integrated ERP &amp; AI Business Management System", styles["title"]), p("A secure multi-branch ERP platform with a grounded AI Intelligence Module", styles["sub"])]
    story += [table(["Proposal detail", "Value"], [
        ["Prepared for", "Client Management Team"], ["Prepared by", "AGI Innovations Private Limited"],
        ["Proposal date", "25 September 2026"], ["Phase 1 delivery", "8 weeks from project kickoff"],
        ["Phase 2 delivery", "3-5 weeks after approved data readiness"], ["Status", "For discussion and scope confirmation"],
    ], [45 * mm, 109 * mm]), Spacer(1, 9 * mm), p("A phased operating platform that enables independent branch operations, central visibility and safe access to the organisation's approved knowledge through AI.", styles["quote"]), Spacer(1, 5 * mm), p("Commercial terms and scope boundaries are detailed in this document. This proposal is valid for 15 days.", styles["note"]), PageBreak()]

    story += [p("1. What the Client Is Building", styles["h1"]), p("The Client requires one cloud-based business management system for multiple branches. Islamabad, Lahore and Karachi will operate within their own data boundaries, while approved management users can see consolidated information. A second phase will add an AI assistant that retrieves answers from approved company material and permitted ERP data.", styles["body"]), p("Business outcomes", styles["h2"])]
    story += bullets(["Single source of truth for sales, stock, service, procurement and finance operations.", "Branch-level data separation, permissions and reporting; no accidental cross-branch data exposure.", "Consolidated management dashboards for authorised leadership users.", "Faster access to policies, records and operational insights through a source-grounded AI assistant.", "A scalable foundation for adding future branches and controlled new capabilities."], styles)
    story += [p("Priority use cases", styles["h2"]), table(["Use case", "Operational effect", "Value hypothesis"], [
        ["Branch operations", "Teams work only within their assigned branch records.", "Risk reduction and cleaner accountability."],
        ["Management reporting", "Authorised users compare branch performance in one workspace.", "Faster oversight and decision-making."],
        ["AI knowledge access", "Staff retrieve approved answers with source references.", "Higher productivity and reduced information search time."],
    ], [43 * mm, 57 * mm, 54 * mm]), Spacer(1, 3 * mm), p("Value posture: the operational benefits above are design objectives, not guaranteed savings. Baseline user volumes, search time and transaction data must be confirmed during discovery before a quantified ROI claim is made.", styles["note"]), PageBreak()]

    story += [p("2. Phase 1 - Multi-Branch ERP Implementation", styles["h1"]), p("Phase 1 converts the current functional ERP experience into a production-ready platform for the initial three branches: Islamabad, Lahore and Karachi.", styles["body"]), table(["Area", "Included Phase 1 deliverables"], [
        ["Secure platform", "Central database, production deployment, authentication, user roles, branch assignment, audit-ready activity logging and backup approach."],
        ["Commercial operations", "CRM, opportunities, quotations, orders, deliveries, invoices, receipts and receivables."],
        ["Inventory & procurement", "Products, warehouses, branch transfers, batches/serials, suppliers, requests, orders and goods receipts."],
        ["Service & oversight", "Equipment, service tickets, maintenance, contracts, reports, approvals and consolidated dashboards."],
        ["Implementation", "Discovery workshops, data templates, agreed initial migration assistance, UAT, training, go-live and handover."],
    ], [46 * mm, 108 * mm]), Spacer(1, 3 * mm), p("Access principle", styles["h2"])]
    story += bullets(["A user is assigned a role and one or more permitted branches.", "Branch users see their assigned records by default; cross-branch access requires explicit authorisation.", "Management visibility is consolidated without giving every user unrestricted access.", "New branches follow the same governed model and are separately configured after Phase 1."], styles)
    story += [PageBreak()]

    story += [p("3. Phase 2 - AI Intelligence Module", styles["h1"]), p("Phase 2 adds a controlled AI assistant inside the ERP. It is not a new model trained from scratch. It uses retrieval-augmented generation (RAG): the assistant retrieves relevant information from approved sources at answer time and responds with references to those sources.", styles["body"]), p("Included capability", styles["h2"])]
    story += bullets(["Chat interface integrated into the ERP workspace.", "Ingestion of approved documents, SOPs, policies and selected ERP records.", "Role- and branch-aware retrieval: a user can retrieve only information they are permitted to access.", "Source references with each answer, plus a safe fallback when the answer is not found in approved material.", "Read-only insights for the first release: search, summaries and operational questions. The assistant will not create, edit or approve business transactions without separately approved workflow controls.", "Quality testing against a jointly agreed set of business questions, plus administrator handover."], styles)
    story += [p("AI guardrails", styles["h2"]), table(["Control", "How it protects the Client"], [
        ["Data boundary", "Branch and role filters are applied before retrieval; sensitive records remain permission controlled."],
        ["Grounded answers", "Responses use approved sources and show references; the system can decline when support is missing."],
        ["Human accountability", "Users remain responsible for decisions; AI output is assistance, not an autonomous approval or financial authority."],
        ["Auditability", "AI queries and operational events can be logged subject to the agreed retention policy."],
    ], [48 * mm, 106 * mm]), Spacer(1, 3 * mm), p("Data readiness dependency: delivery assumes the Client supplies approved, organised documents and identifies data owners. Extensive legacy-data cleansing, multilingual corpus preparation and self-hosted/private model infrastructure are not included unless separately scoped.", styles["note"]), PageBreak()]

    story += [p("4. Delivery & Governance", styles["h1"]), table(["Phase", "Timing", "Decision-useful outcome"], [
        ["ERP discovery & blueprint", "Week 1", "Confirm branches, roles, workflows, data, reports and acceptance criteria."],
        ["ERP build, testing & go-live", "Weeks 2-8", "Production-ready initial branches, UAT, training and launch support."],
        ["AI discovery & data readiness", "Week 1 of Phase 2", "Approve sources, permissions, test questions, guardrails and success criteria."],
        ["AI build, evaluation & launch", "Weeks 2-5 of Phase 2", "Grounded assistant, access controls, evaluation, training and controlled release."],
    ], [51 * mm, 31 * mm, 72 * mm]), Spacer(1, 3 * mm), p("Client responsibilities", styles["h2"])]
    story += bullets(["Nominate a project owner, branch key users and an AI data owner.", "Provide timely process decisions, clean agreed data/templates and approved source documents.", "Review demonstrations, complete user acceptance testing and approve milestone acceptance within five business days.", "Confirm the permissions matrix and business questions used to evaluate AI answers."], styles)
    story += [p("Change control", styles["h2"]), p("Any scope outside this proposal will be documented with impact on delivery dates and cost before work begins. This protects the Phase 1 and Phase 2 timelines and avoids uncontrolled changes during implementation.", styles["body"]), PageBreak()]

    story += [p("5. Commercial Proposal", styles["h1"]), p("The engagement is intentionally phased so the ERP foundation and its branch controls are established before AI is given access to approved business information.", styles["body"]), table(["Investment", "Fee", "Included"], [
        ["Phase 1: Multi-Branch ERP", "PKR 1,200,000", "Initial Islamabad, Lahore and Karachi rollout; implementation scope in Section 2."],
        ["Phase 2: AI Intelligence Module", "USD 2,500 / PKR equivalent at invoice date", "Controlled RAG assistant, ERP integration, branch/role-aware retrieval, evaluation and launch."],
        ["Taxes", "Excluded", "Applicable taxes, if any, are charged as required by law."],
        ["Recurring services", "Excluded", "LLM/API use, cloud hosting, vector database, domain and third-party service charges."],
    ], [54 * mm, 48 * mm, 52 * mm]), Spacer(1, 4 * mm), p("Payment schedule", styles["h2"]), table(["Milestone", "Phase 1 ERP", "Phase 2 AI"], [
        ["Kickoff & discovery approval", "40% - PKR 480,000", "50% - USD 1,250"],
        ["UAT-ready solution", "40% - PKR 480,000", "30% - USD 750"],
        ["Go-live & handover", "20% - PKR 240,000", "20% - USD 500"],
    ], [60 * mm, 47 * mm, 47 * mm]), Spacer(1, 3 * mm), p("Commercial note", styles["h2"]), p("Phase 2 pricing applies to the defined v1 assistant. Autonomous write-actions, complex external integrations, private/on-premise model hosting, extensive data remediation, custom mobile AI apps and advanced analytics are subject to separate scope and quotation.", styles["body"]), PageBreak()]

    story += [p("6. Assumptions, Support & Next Decision", styles["h1"]), p("Included support", styles["h2"]), p("Each phase includes 30 calendar days of remote post-go-live warranty support for defects within the delivered scope. Enhancements, additional support coverage and recurring infrastructure services may be covered by a separate support agreement.", styles["body"]), p("Not included unless separately agreed", styles["h2"])]
    story += bullets(["Additional branches after Islamabad, Lahore and Karachi; indicative setup fee: PKR 125,000 per branch, subject to confirmation.", "FBR/e-invoicing, banking, payment gateway, biometric, GPS, SMS/WhatsApp or third-party integrations.", "Hardware purchases, major historical transaction migration, extensive source-document cleansing or bespoke reports.", "Guaranteed AI accuracy, automated financial/operational approvals, or replacement of human review."], styles)
    story += [p("Recommended decision", styles["h2"]), p("Approve Phase 1, nominate the Client project owner and branch key users, and schedule the discovery workshop. Phase 2 should be approved after the data sources, access matrix and AI evaluation questions are validated during the ERP rollout.", styles["body"]), Spacer(1, 9 * mm), table(["For AGI Innovations Private Limited", "For the Client"], [["\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________", "\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________"]], [77 * mm, 77 * mm])]
    ProposalDoc(str(PDF_OUT)).build(story)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), color)
    props.append(item)


def cell(cell, text, bold=False, white=False):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.bold = bold
    run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(32, 33, 36)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def word_table(doc, headers, rows):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for i, header in enumerate(headers):
        shade(tbl.rows[0].cells[i], "32156F")
        cell(tbl.rows[0].cells[i], header, True, True)
    for r, row in enumerate(rows):
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            shade(cells[i], "F1EDF8" if r % 2 else "FFFFFF")
            cell(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_docx():
    assets = OUT_DIR / "proposal_assets"
    assets.mkdir(exist_ok=True)
    head, foot = assets / "agi_header.png", assets / "agi_footer.png"
    image = Image.open(LETTERHEAD)
    image.crop((0, 0, image.width, 238)).save(head)
    image.crop((0, 1424, image.width, image.height)).save(foot)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2), Cm(4.8), Cm(2.1)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    for name, size, color in [("Normal", 8.5, (32, 33, 36)), ("Title", 21, (50, 21, 111)), ("Heading 1", 13, (50, 21, 111)), ("Heading 2", 10, (50, 21, 111))]:
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.color.rgb = "Arial", Pt(size), RGBColor(*color)
        if name != "Normal": st.font.bold = True
    hdr = sec.header.paragraphs[0]; hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER; hdr.add_run().add_picture(str(head), width=Cm(21))
    ftr = sec.footer.paragraphs[0]; ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER; ftr.add_run().add_picture(str(foot), width=Cm(21))

    def para(text="", style=None, center=False, bold=False, color=None, size=None):
        q = doc.add_paragraph(style=style); q.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        r = q.add_run(text); r.bold = bold
        if color: r.font.color.rgb = RGBColor(*color)
        if size: r.font.size = Pt(size)
        return q
    def list_item(text):
        q = doc.add_paragraph(style="List Bullet"); q.paragraph_format.space_after = Pt(2); q.add_run(text)

    para("CONFIDENTIAL BUSINESS PROPOSAL", center=True, bold=True, color=(85, 48, 138), size=8.5)
    para("Integrated ERP & AI Business Management System", "Title", True)
    para("A secure multi-branch ERP platform with a grounded AI Intelligence Module", center=True, size=10.5)
    word_table(doc, ["Proposal detail", "Value"], [["Prepared for", "Client Management Team"], ["Prepared by", "AGI Innovations Private Limited"], ["Proposal date", "25 September 2026"], ["Phase 1 delivery", "8 weeks from project kickoff"], ["Phase 2 delivery", "3-5 weeks after approved data readiness"]])
    para("A phased operating platform that enables independent branch operations, central visibility and safe access to approved business knowledge through AI.", center=True, bold=True, color=(50, 21, 111), size=10)
    doc.add_page_break()
    sections = [
        ("1. What the Client Is Building", "The Client requires one cloud-based business management system for multiple branches. Islamabad, Lahore and Karachi will operate within their own data boundaries, while approved management users can see consolidated information. A second phase will add an AI assistant that retrieves answers from approved company material and permitted ERP data.", ["Single source of truth for sales, stock, service, procurement and finance operations.", "Branch-level data separation, permissions and reporting.", "Consolidated management dashboards for authorised leadership users.", "Faster access to approved knowledge through a source-grounded AI assistant."]),
        ("2. Phase 1 - Multi-Branch ERP", "Phase 1 converts the current ERP experience into a production-ready platform for Islamabad, Lahore and Karachi. It includes central data persistence, secure authentication, role/branch controls, sales, inventory, procurement, service, finance, reports, migration assistance, UAT, training and launch support.", ["Branch users see only assigned records by default.", "Management can receive consolidated visibility without unrestricted user access.", "Additional branches follow the same governed model and are separately configured."]),
        ("3. Phase 2 - AI Intelligence Module", "The AI module is a controlled retrieval-augmented generation (RAG) assistant, not a new model trained from scratch. It retrieves approved documents and permitted ERP data at answer time and provides source references.", ["ERP-integrated chat interface and approved document/SOP ingestion.", "Role- and branch-aware retrieval with safe fallback when support is missing.", "Read-only v1 search, summaries and operational insights; no autonomous transaction edits or approvals.", "Joint evaluation against agreed business questions and administrator handover."]),
        ("4. Delivery & Governance", "Phase 1 runs for eight weeks. Phase 2 requires a further three to five weeks after approved data readiness. The Client provides a project owner, branch key users, an AI data owner, approved sources, a permissions matrix and timely UAT feedback.", ["ERP discovery and blueprint: Week 1; build, UAT and go-live: Weeks 2-8.", "AI data/guardrail workshop: Phase 2 Week 1; build, evaluation and controlled launch: Weeks 2-5.", "Scope changes require written approval with delivery and cost impact."]),
    ]
    for number, text, points in sections:
        para(number, "Heading 1"); para(text)
        for item in points: list_item(item)
        doc.add_page_break()
    para("5. Commercial Proposal", "Heading 1")
    word_table(doc, ["Investment", "Fee", "Included"], [["Phase 1: Multi-Branch ERP", "PKR 1,200,000", "Initial Islamabad, Lahore and Karachi rollout."], ["Phase 2: AI Intelligence Module", "USD 2,500 / PKR equivalent at invoice date", "Controlled RAG assistant, ERP integration, access controls, evaluation and launch."], ["Taxes", "Excluded", "Applicable taxes charged as required by law."], ["Recurring services", "Excluded", "LLM/API, hosting, vector database and third-party service charges."]])
    para("Payment Schedule", "Heading 2")
    word_table(doc, ["Milestone", "Phase 1 ERP", "Phase 2 AI"], [["Kickoff & discovery", "40% - PKR 480,000", "50% - USD 1,250"], ["UAT-ready solution", "40% - PKR 480,000", "30% - USD 750"], ["Go-live & handover", "20% - PKR 240,000", "20% - USD 500"]])
    para("6. Assumptions, Support & Next Decision", "Heading 1")
    para("Each phase includes 30 calendar days of remote warranty support for delivered-scope defects. Additional branches, third-party integrations, extensive historical migration, private/on-prem model hosting, autonomous AI actions and ongoing services are separately scoped.")
    para("Recommended Decision", "Heading 2")
    para("Approve Phase 1, nominate the Client project owner and branch key users, and schedule the discovery workshop. Approve Phase 2 after the data sources, access matrix and AI evaluation questions are validated.")
    word_table(doc, ["For AGI Innovations Private Limited", "For the Client"], [["\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________", "\n\nAuthorised Signatory\nName: __________________________\nDate: __________________________"]])
    doc.save(DOCX_OUT)


if __name__ == "__main__":
    build_pdf()
    build_docx()
    print(PDF_OUT)
    print(DOCX_OUT)
